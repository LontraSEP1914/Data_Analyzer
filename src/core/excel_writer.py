import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, numbers
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

class ExcelWriter:
    """
    Classe responsável por escrever e formatar o arquivo Excel final.
    """
    def __init__(self):
        pass

    def write_to_excel(self, file_path: str, details_df: pl.DataFrame, summary_data: dict):
        """
        Escreve o DataFrame de detalhes e o resumo em um arquivo Excel com duas sheets.

        Args:
            file_path (str): O caminho e nome do arquivo de saída.
            details_df (pl.DataFrame): O DataFrame com o resultado detalhado da operação.
            summary_data (dict): Dicionário com os dados resumidos da operação.
        """
        print(f"Gerando o arquivo Excel em: {file_path}")
        wb = Workbook()

        # Sheet de Detalhes
        ws_details = wb.active
        ws_details.title = "Detalhes"
        
        rows = dataframe_to_rows(details_df.to_pandas(), index=False, header=True)
        for r_idx, row in enumerate(rows, 1):
            ws_details.append(row)
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
        for cell in ws_details[1]:
            cell.font = header_font
            cell.fill = header_fill

        for col_idx, col_name in enumerate(details_df.columns, 1):
            if col_name.startswith('DIF_%_'):
                column_letter = get_column_letter(col_idx)
                for cell in ws_details[column_letter][1:]:
                    cell.number_format = numbers.BUILTIN_FORMATS[10]

        ws_details.freeze_panes = 'A2'
        ws_details.sheet_view.showGridLines = False
        ws_details.sheet_view.zoomScale = 70

        # Sheet de Resumo
        ws_summary = wb.create_sheet("Resumo")
        
        # AQUI FOI ADICIONADO A FORMATAÇÃO PARA A SHEET DE RESUMO
        ws_summary.sheet_view.showGridLines = False
        ws_summary.sheet_view.zoomScale = 70
        
        summary_header = ['Descricao', 'Valor']
        ws_summary.append(summary_header)
        for cell in ws_summary[1]:
            cell.font = header_font
            cell.fill = header_fill
            
        for key, value in summary_data.items():
            row = [key, value]
            ws_summary.append(row)
            
            if 'Diferença %' in key or '%' in key:
                ws_summary.cell(row=ws_summary.max_row, column=2).number_format = numbers.BUILTIN_FORMATS[10]
        
        for col in ws_summary.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws_summary.column_dimensions[col[0].column_letter].width = adjusted_width

        try:
            wb.save(file_path)
            print("Arquivo Excel gerado com sucesso!")
        except Exception as e:
            print(f"Erro ao salvar o arquivo Excel: {e}")
            raise

if __name__ == '__main__':
    # Exemplo de uso para testes manuais
    from polars_processor import DataProcessor

    processor = DataProcessor()
    
    # Simula o cruzamento
    df_a_cruz = pl.DataFrame({"ID": [1, 2, 3], "Produto": ["A", "B", "C"]})
    df_b_cruz = pl.DataFrame({"ID": [2, 3, 4], "Preco": [10, 20, 30]})
    cruzamento_df = processor.cruzamento(df_a_cruz, df_b_cruz, on_columns=["ID"])
    
    # Dados de resumo para o cruzamento
    cruzamento_summary = {
        'Linhas processadas': len(df_a_cruz),
        'Linhas com correspondência': len(cruzamento_df.filter(~pl.col("Preco").is_null()))
    }

    excel_writer = ExcelWriter()
    excel_writer.write_to_excel('cruzamento_final.xlsx', cruzamento_df, cruzamento_summary)

    # Simula o confronto
    df_a_confr = pl.DataFrame({"ID": [1, 2, 3], "Vendas": [100, 200, 150]})
    df_b_confr = pl.DataFrame({"ID": [1, 2, 3], "Vendas": [120, 180, 150]})
    confronto_df = processor.confronto(df_a_confr, df_b_confr, on_columns=["ID"], value_columns=["Vendas"])
    
    # Dados de resumo para o confronto
    confronto_summary = {
        'Linhas processadas': len(confronto_df),
        'Total de Diferença': confronto_df.select(pl.col('DIF_Vendas').sum()).item(),
        'Média de Diferença %': confronto_df.select(pl.col('DIF_%_Vendas').mean()).item()
    }

    excel_writer.write_to_excel('confronto_final.xlsx', confronto_df, confronto_summary)